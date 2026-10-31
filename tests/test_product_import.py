import io

import httpx
from openpyxl import Workbook, load_workbook

from app.routers import products as products_router
from tests.conftest import login as _login

_tokens: dict = {}


async def login(client, phone):
    if phone not in _tokens:
        _tokens[phone] = await _login(client, phone)
    return _tokens[phone]


def _xlsx(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_template_preview_and_import(client, monkeypatch):
    h = await login(client, "09127777777")
    sid = (await client.post("/stores", json={"name": "اکسل"}, headers=h)).json()["id"]

    tpl = await client.get(f"/stores/{sid}/products/import-template", headers=h)
    assert tpl.status_code == 200
    header = [c.value for c in next(load_workbook(io.BytesIO(tpl.content)).worksheets[0].iter_rows(max_row=1))]
    assert header[0] == "عنوان" and "کد کالا" in header
    csv_tpl = await client.get(f"/stores/{sid}/products/import-template?format=csv", headers=h)
    assert csv_tpl.content.decode("utf-8-sig").startswith("عنوان,")

    # an existing product with a SKU gets updated instead of duplicated
    await client.post(f"/stores/{sid}/products", headers=h, json={"title": "ماگ قدیمی", "sku": "MG-1", "price": 100})

    data = _xlsx([["عنوان", "قیمت", "موجودی", "کد کالا", "برچسب‌ها", "لینک عکس‌ها"],
                  ["هودی مشکی", "۸۵۰٬۰۰۰ تومان", 12, "HD-1", "هودی، زمستانه", "https://img.test/a.png"],
                  ["ماگ سفید", 120000, None, "MG-1", None, None],
                  [None, 5000, 1, None, None, None],
                  ["کلاه", "ارزان", 1, None, None, None]])
    r = await client.post(f"/stores/{sid}/products/import/preview", headers=h,
                          files={"file": ("p.xlsx", data, "application/octet-stream")})
    assert r.status_code == 200, r.text
    pv = r.json()
    assert pv["total"] == 4 and pv["valid"] == 2 and pv["create"] == 1 and pv["update"] == 1
    first = pv["rows"][0]["data"]
    assert first["price"] == 850000 and first["stock_qty"] == 12 and first["tags"] == ["هودی", "زمستانه"]
    assert pv["rows"][2]["errors"] == ["عنوان خالی است."] and "عدد معتبر نیست" in pv["rows"][3]["errors"][0]

    png = b"\x89PNG\r\n\x1a\n" + b"0" * 32
    real = httpx.AsyncClient
    monkeypatch.setattr(products_router.httpx, "AsyncClient", lambda *a, **kw: real(
        *a, transport=httpx.MockTransport(lambda req: httpx.Response(200, content=png, headers={"content-type": "image/png"})), **kw))
    items = [row["data"] for row in pv["rows"] if not row["errors"]]
    res = (await client.post(f"/stores/{sid}/products/import", headers=h, json={"items": items})).json()
    assert res == {"created": 1, "updated": 1, "image_failed": 0}

    prods = (await client.get(f"/stores/{sid}/products", headers=h)).json()["items"]
    by_sku = {p["sku"]: p for p in prods}
    assert len(prods) == 2 and by_sku["MG-1"]["price"] == 120000 and by_sku["MG-1"]["title"] == "ماگ سفید"
    assert by_sku["HD-1"]["images"][0]["url"].startswith("/media/")

    # wrong files are explained
    r = await client.post(f"/stores/{sid}/products/import/preview", headers=h, files={"file": ("p.xls", b"x", "application/octet-stream")})
    assert r.status_code == 422 and "xlsx" in r.json()["detail"]["message"]
    r = await client.post(f"/stores/{sid}/products/import/preview", headers=h,
                          files={"file": ("p.csv", "name2,price\nx,1\n".encode(), "text/csv")})
    assert r.status_code == 422 and "عنوان" in r.json()["detail"]["message"]
    # CSV with English headers works too
    r = await client.post(f"/stores/{sid}/products/import/preview", headers=h,
                          files={"file": ("p.csv", "title,price,stock\nشال,90000,3\n".encode(), "text/csv")})
    assert r.json()["valid"] == 1 and r.json()["rows"][0]["data"]["price"] == 90000
