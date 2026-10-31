"""Products from a spreadsheet (Excel .xlsx or CSV): template, parsing and per-row validation.

Headers may be Persian (as in the template) or English; column order doesn't matter. Rows are checked
before anything is saved, so the seller sees every problem in the preview first."""

import csv
import io
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

MAX_ROWS = 1000
MAX_BYTES = 5 * 1024 * 1024

# key, Persian header, accepted aliases, example
COLUMNS: list[tuple[str, str, tuple[str, ...], str]] = [
    ("title", "عنوان", ("نام", "نام محصول", "عنوان محصول", "title", "name", "product"), "هودی مشکی دورس"),
    ("price", "قیمت (تومان)", ("قیمت", "قیمت فروش", "price"), "850000"),
    ("stock_qty", "موجودی", ("تعداد", "stock", "qty", "quantity", "stock_qty"), "12"),
    ("sku", "کد کالا", ("کد", "sku", "code", "شناسه"), "HD-001"),
    ("compare_at_price", "قیمت قبل از تخفیف", ("قیمت قبل", "compare_at_price", "old_price"), "990000"),
    ("brand", "برند", ("brand",), ""),
    ("category_hint", "دسته", ("دسته‌بندی", "دسته بندی", "category"), "پوشاک"),
    ("description", "توضیحات", ("توضیح", "description"), "جنس پنبه، سایز L و XL"),
    ("tags", "برچسب‌ها", ("برچسب", "تگ", "tags"), "هودی، زمستانه"),
    ("images", "لینک عکس‌ها", ("لینک عکس", "عکس", "تصویر", "images", "image", "image_url"), ""),
]
REQUIRED = {"title"}
NUMERIC = {"price", "stock_qty", "compare_at_price"}

_FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def _norm_header(h: Any) -> str:
    s = str(h or "").strip().lower().replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
    return re.sub(r"\s+", " ", s)


_HEADER_MAP: dict[str, str] = {}
for _key, _fa, _aliases, _ in COLUMNS:
    for _name in (_fa, *_aliases):
        _HEADER_MAP[_norm_header(_name)] = _key


class ImportError_(ValueError):
    """A file-level problem, worded for the seller."""


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "محصولات"
    ws.sheet_view.rightToLeft = True
    ws.append([fa for _, fa, _, _ in COLUMNS])
    ws.append([ex for *_, ex in COLUMNS])
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="D1FAE5")
    for c in ws[1]:
        c.font, c.fill = bold, fill
    for i, (key, fa, _, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 40 if key in ("description", "images") else max(14, len(fa) + 6)
    help_ws = wb.create_sheet("راهنما")
    help_ws.sheet_view.rightToLeft = True
    for line in ("فقط ستون «عنوان» اجباری است؛ بقیه اختیاری‌اند.",
                 "قیمت‌ها به تومان و بدون جداکننده یا با جداکننده؛ عدد فارسی هم قبول است.",
                 "اگر «کد کالا» با محصولی که از قبل داری یکی باشد، همان محصول به‌روز می‌شود.",
                 "برچسب‌ها و لینک عکس‌ها را با ویرگول (، یا ,) جدا کن. حداکثر ۵ عکس برای هر محصول.",
                 "ردیف نمونه را قبل از بارگذاری پاک کن."):
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_csv() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([fa for _, fa, _, _ in COLUMNS])
    w.writerow([ex for *_, ex in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _rows_from_csv(data: bytes) -> list[list[Any]]:
    for enc in ("utf-8-sig", "cp1256"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportError_("فایل CSV خوانده نشد؛ آن را با رمزگذاری UTF-8 ذخیره کن.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [r for r in csv.reader(io.StringIO(text), dialect)]


def _rows_from_xlsx(data: bytes) -> list[list[Any]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any broken workbook gets the same plain message
        raise ImportError_("فایل اکسل باز نشد؛ مطمئن شو فرمتش xlsx است.") from exc
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def read_rows(filename: str, data: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    """(recognized column keys, [{"row": n, **raw by key}]). Raises ImportError_ for file-level problems."""
    if len(data) > MAX_BYTES:
        raise ImportError_("حجم فایل بیشتر از ۵ مگابایت است.")
    name = (filename or "").lower()
    if name.endswith(".xls"):
        raise ImportError_("فرمت قدیمی xls پشتیبانی نمی‌شود؛ فایل را در اکسل با فرمت xlsx یا CSV ذخیره کن.")
    rows = _rows_from_xlsx(data) if name.endswith((".xlsx", ".xlsm")) else _rows_from_csv(data)
    rows = [r for r in rows if any(str(c or "").strip() for c in r)]
    if not rows:
        raise ImportError_("فایل خالی است.")
    keys = [_HEADER_MAP.get(_norm_header(h)) for h in rows[0]]
    if "title" not in keys:
        raise ImportError_("ستون «عنوان» در ردیف اول پیدا نشد؛ از فایل نمونه استفاده کن.")
    body = rows[1:]
    if len(body) > MAX_ROWS:
        raise ImportError_(f"هر بار حداکثر {MAX_ROWS} محصول؛ فایل را چند تکه کن.")
    out = []
    for i, r in enumerate(body, start=2):
        item: dict[str, Any] = {"row": i}
        for key, cell in zip(keys, r):
            if key and cell is not None and str(cell).strip():
                item[key] = cell
        out.append(item)
    return [k for k in dict.fromkeys(keys) if k], out


def _number(v: Any) -> int | None:
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).translate(_FA_DIGITS).replace("٬", "").replace(",", "").replace("،", "").strip()
    s = re.sub(r"(تومان|ریال|toman)", "", s, flags=re.I).strip()
    try:
        return int(float(s))
    except ValueError:
        return None


def _split(v: Any) -> list[str]:
    return [x.strip() for x in re.split(r"[،,\n]", str(v)) if x.strip()]


def clean_row(raw: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    """Validated product fields plus Persian problems for this row (empty when it can be saved)."""
    data: dict[str, Any] = {}
    errors: list[str] = []
    title = str(raw.get("title", "")).strip()
    if not title:
        errors.append("عنوان خالی است.")
    data["title"] = title[:300]
    for key in NUMERIC:
        if key in raw:
            n = _number(raw[key])
            if n is None or n < 0:
                label = next(fa for k, fa, _, _ in COLUMNS if k == key)
                errors.append(f"«{label}» عدد معتبر نیست: {raw[key]}")
            else:
                data[key] = n
    for key in ("sku", "brand", "category_hint", "description"):
        if key in raw:
            v = raw[key]
            data[key] = str(int(v) if isinstance(v, float) and v.is_integer() else v).strip()
    if "tags" in raw:
        data["tags"] = [t.lstrip("#") for t in _split(raw["tags"])][:8]
    if "images" in raw:
        urls = _split(raw["images"])
        bad = [u for u in urls if not u.startswith(("http://", "https://"))]
        if bad:
            errors.append("لینک عکس باید با http یا https شروع شود.")
        data["images"] = [u for u in urls if u not in bad][:5]
    return data, errors
